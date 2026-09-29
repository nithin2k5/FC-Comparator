"""CSV / Excel export of inspection history and the daily report."""

from __future__ import annotations

import csv
from datetime import date, datetime
from pathlib import Path

from .database import DailyReport, InspectionStore

INSPECTION_COLUMNS = [
    "id", "ts", "station", "operator_id", "part_number", "mode", "result", "ng_count",
    "duration_ms", "detector", "placement_ok", "placement_shift", "placement_message", "extras_count", "error",
    "image_path",
]
POSITION_COLUMNS = ["inspection_id", "cable", "row", "expected", "found", "confidence", "ok", "reason"]


def _collect(store: InspectionStore, start: datetime | None, end: datetime | None, **filters):
    rows = list(reversed(store.inspections(start, end, **filters)))  # oldest first for export
    positions = store.positions([r["id"] for r in rows])
    mismatches: dict[int, list[str]] = {}
    for p in positions:
        if not p["ok"]:
            mismatches.setdefault(p["inspection_id"], []).append(
                f"C{p['cable']}R{p['row']}:{p['found']}!={p['expected']}({p['confidence']:.2f})"
            )
    return rows, positions, mismatches


def export_csv(
    store: InspectionStore, path: str | Path, start: datetime | None = None, end: datetime | None = None, **filters
) -> tuple[Path, Path]:
    """Write ``<path>`` (one row per inspection) and ``<path stem>_positions.csv``."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    rows, positions, mismatches = _collect(store, start, end, **filters)
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:  # BOM so Excel opens UTF-8 correctly
        w = csv.writer(fh)
        w.writerow(INSPECTION_COLUMNS + ["mismatches"])
        for r in rows:
            w.writerow([r[c] for c in INSPECTION_COLUMNS] + [" ".join(mismatches.get(r["id"], []))])
    pos_path = path.with_name(path.stem + "_positions.csv")
    with open(pos_path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(POSITION_COLUMNS)
        for p in positions:
            w.writerow([p[c] for c in POSITION_COLUMNS])
    return path, pos_path


def export_excel(
    store: InspectionStore,
    path: str | Path,
    start: datetime | None = None,
    end: datetime | None = None,
    report_days: list[date] | None = None,
    **filters,
) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("Excel export needs openpyxl: pip install openpyxl") from exc

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    rows, positions, mismatches = _collect(store, start, end, **filters)
    wb = Workbook()
    bold = Font(bold=True)
    red = PatternFill("solid", fgColor="FFC7CE")

    ws = wb.active
    ws.title = "Inspections"
    ws.append(INSPECTION_COLUMNS + ["mismatches"])
    for r in rows:
        ws.append([r[c] for c in INSPECTION_COLUMNS] + [" ".join(mismatches.get(r["id"], []))])
        if r["result"] == "NG":
            for cell in ws[ws.max_row]:
                cell.fill = red

    wp = wb.create_sheet("Positions")
    wp.append(POSITION_COLUMNS)
    for p in positions:
        wp.append([p[c] for c in POSITION_COLUMNS])

    days = report_days or sorted({datetime.fromisoformat(r["ts"]).date() for r in rows})
    wr = wb.create_sheet("Daily report")
    wr.append(["Day", "Inspected", "OK", "NG", "NG rate", "Top failing positions", "Top failures"])
    for d in days:
        rep = store.daily_report(d)
        wr.append([
            d.isoformat(), rep.total, rep.ok, rep.ng, round(rep.ng_rate, 4),
            ", ".join(f"C{c}R{r} x{n}" for c, r, n in rep.failing_positions),
            ", ".join(f"{f} for {e} x{n}" for e, f, n in rep.failure_kinds),
        ])

    for sheet in (ws, wp, wr):
        for cell in sheet[1]:
            cell.font = bold
        sheet.freeze_panes = "A2"
        for col in sheet.columns:
            width = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            sheet.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
    wb.save(path)
    return path


def export_daily_report_text(report: DailyReport, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(report.as_text() + "\n", encoding="utf-8")
    return path
